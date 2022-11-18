import docx
import os
from docx import Document
import xlrd
from xlutils.copy import copy

def read_excel(a,cout):
    b = []
    text = []
    xlrdBook = xlrd.open_workbook('./train/222.xls')
    wb = copy(xlrdBook)
    ws = wb.get_sheet(0)
    tablies = xlrdBook.sheet_by_name('Sheet1')
    for i in range(0, tablies.ncols):
        b.append(str(tablies.cell(0, i).value))
    for i in b:
        for j in range(0, len(a)):
            if i==a[j]:
                text.append(a[j+1])
    print(text)
    for i in range(0, tablies.ncols):
        ws.write(cout, i, label=text[i])
    wb.save('./train/222.xls')
def read_doc(path):
    dirs = os.listdir(path)
    for cout in range(0,len(dirs)):
        doc_path = os.path.join(path, dirs[cout])
        doc = Document(doc_path)
        table = doc.tables[0]
        author_list_elems = []
        a = []
        for row in table.rows:
            for cell in row.cells:
                text = ''.join(cell.text.replace('\n','').replace(' ',''))
                author_list_elems.append(text)
        author_list_elems = sorted(set(author_list_elems), key=author_list_elems.index)
        for text in author_list_elems:
            a.append(text)
            if text == '全日制':
                break
        for i in range(0,len(a)):
            if a[i] == '一志愿报考专业代码':
                code = a[i+1]
            if a[i] == '一志愿报考专业名称':
                name = a[i+1]
            if a[i] == '专业二科目名称':
                a.insert(i+2, '计算机专业课程综合成绩')
            if a[i] == '代码':
                if type(a[i+1]) is not int:
                    a.insert(i+1, code)
            if a[i] == '名称':
                if type(a[i + 1]) is not int:
                    a.insert(i+1, name)
        print(a)
        read_excel(a, cout+1)
if __name__ == '__main__':
    path = './test'
    read_doc(path)